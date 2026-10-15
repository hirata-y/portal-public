#!/usr/bin/env python3
"""Scan real build, repeat the same scanner on fault-injected temporary copies."""
import gzip
import json,hashlib,tempfile,shutil,subprocess,sys,zipfile
from pathlib import Path
import fitz
from pptx import Presentation
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1];D=R/'astro_candidate/dist';E=R/'evidence/vendor_actual';C=json.loads((R/'docs/SOURCE_PIN.json').read_text())['canary'].encode();checks=[]
def ck(id,ok,actual):checks.append({'id':id,'status':'PASS' if ok else 'FAIL','actual':actual})
from privacy_scan import scan as scan_root
scan=lambda root:scan_root(root,C)
ck('REAL_DIST_SAME_SCANNER_CLEAN',scan(D)==[],scan(D))
with tempfile.TemporaryDirectory() as t:
 for name,rel,data in [('html_canary','index.html',C),('index_canary','pagefind/injected.pf_fragment',gzip.compress(b'pagefind_dcd'+C)),('private_index','members-index.json',b'{}'),('private_route','members/canary.html',b'x')]:
  dest=Path(t)/name;shutil.copytree(D,dest);p=dest/rel;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(data);found=scan(dest);ck('MUTATION_'+name,any(rel in x for x in found),found)
 # Generated real Pagefind compressed index is decompressed and queried by the browser probe separately.
# Input ZIP preservation and source subtree byte identity.
z=zipfile.ZipFile(R/'source/P3_B1_VENDOR_READINESS_v0.2.1_20261015.zip');prefix=z.namelist()[0].split('/')[0]+'/'
for path in sorted((R/'source/p3_a_site').rglob('*')):
 if path.is_file():
  rel=path.relative_to(R).as_posix();ck('FIXED_SOURCE_'+str(path.relative_to(R/'source/p3_a_site')),path.read_bytes()==z.read(prefix+rel),hashlib.sha256(path.read_bytes()).hexdigest())
f=D/'fixtures';out=E/'formats';out.mkdir(exist_ok=True)
pdf=fitz.open(f/'synthetic_format.pdf');ck('PDF_PARSE_CONTENT',len(pdf)==1 and bool(pdf[0].get_text().strip()),{'pages':len(pdf),'text':pdf[0].get_text()});pdf[0].get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(str(out/'pdf_page1.png'))
ppt=Presentation(f/'synthetic_format.pptx');texts=[s.text for sl in ppt.slides for s in sl.shapes if s.has_text_frame];ck('PPTX_PARSE_CONTENT',len(ppt.slides)==1 and bool(texts),texts)
wb=load_workbook(f/'synthetic_format.xlsx',data_only=False);cells={s.title:list(s.values) for s in wb};ck('XLSX_PARSE_CONTENT',bool(cells) and any(any(v is not None for v in row) for s in wb for row in s.values),cells)
ck('HTML_ACTUAL_CONTENT','合成' in (f/'synthetic_format.html').read_text() or 'Synthetic' in (f/'synthetic_format.html').read_text(),(f/'synthetic_format.html').read_text())
report={'role':'A Producer self-check','checks':checks,'summary':{s:sum(c['status']==s for c in checks) for s in ['PASS','FAIL']},'limits':['Synthetic public/member output segregation only; not real authorization','PDF rendered; PPTX/XLSX parsed and rendered separately with LibreOffice; Microsoft Office desktop unverified']};(E/'security_formats.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str));print(json.dumps(report,ensure_ascii=False,indent=2,default=str));sys.exit(0 if report['summary']['FAIL']==0 else 2)
